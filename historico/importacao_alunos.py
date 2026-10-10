import re
import unicodedata
import zipfile
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from django.db import transaction
from django.db.models import Max

from .models import Aluno


MAX_LINHAS = 20000
MAX_COLUNAS = 80
MAX_XLSX_DESCOMPACTADO = 120 * 1024 * 1024
MAX_XLSX_ARQUIVOS_INTERNOS = 5000

CABECALHOS = {
    "codigo": {
        "codigo", "codigo aluno", "cod aluno", "cod do aluno",
        "codigo do aluno", "cod", "registro aluno",
    },
    "nome": {
        "nome", "nome aluno", "nome do aluno", "aluno",
    },
    "matricula": {
        "matricula", "matricula aluno", "matricula do aluno",
    },
    "cpf": {"cpf", "cpf aluno", "cpf do aluno"},
    "pai": {
        "pai", "nome pai", "nome do pai", "filiacao pai",
        "filiacao 1", "filiacao1",
    },
    "mae": {
        "mae", "nome mae", "nome da mae", "filiacao mae",
        "filiacao 2", "filiacao2",
    },
    "nascimento": {
        "nascimento", "data nascimento", "data de nascimento",
        "dt nascimento", "dt nasc", "data nasc",
    },
    "naturalidade": {"naturalidade", "cidade nascimento", "natural"},
    "nacionalidade": {
        "nacionalidade", "nacionalidad", "nacional",
    },
    "sexo": {"sexo", "genero", "sexo aluno"},
    "uf": {
        "uf", "uf nascimento", "uf de nascimento", "estado nascimento",
        "estado de nascimento",
    },
    "identidade": {
        "identidade", "rg", "carteira identidade",
        "carteira de identidade",
    },
    "orgao_expedidor": {
        "orgao expedidor", "orgao expedidor estado", "orgao rg",
        "orgao emissor",
    },
    "curso": {"curso", "nivel ensino", "nivel de ensino"},
}

ROTULOS = {
    "codigo": "Código",
    "nome": "Nome",
    "matricula": "Matrícula",
    "cpf": "CPF",
    "pai": "Pai",
    "mae": "Mãe",
    "nascimento": "Nascimento",
    "naturalidade": "Naturalidade",
    "nacionalidade": "Nacionalidade",
    "sexo": "Sexo",
    "uf": "UF",
    "identidade": "Identidade",
    "orgao_expedidor": "Órgão expedidor",
    "curso": "Curso",
}

CAMPOS_ATUALIZAVEIS = (
    "nome",
    "matricula",
    "cpf",
    "curso",
    "identidade",
    "orgao_expedidor",
    "uf",
    "pai",
    "mae",
    "nascimento",
    "naturalidade",
    "nacionalidade",
    "sexo",
)


class ErroImportacao(ValueError):
    pass


def _sem_acentos(valor):
    return "".join(
        char
        for char in unicodedata.normalize("NFKD", str(valor or ""))
        if not unicodedata.combining(char)
    )


def _normalizar_texto(valor):
    texto = _sem_acentos(valor).upper()
    return re.sub(r"\s+", " ", texto).strip()


def _normalizar_cabecalho(valor):
    texto = _sem_acentos(valor).lower()
    texto = re.sub(r"[^a-z0-9]+", " ", texto)
    return re.sub(r"\s+", " ", texto).strip()


def _mapa_aliases():
    resultado = {}
    for campo, aliases in CABECALHOS.items():
        for alias in aliases:
            resultado[_normalizar_cabecalho(alias)] = campo
    return resultado


ALIASES = _mapa_aliases()


def _texto_celula(valor):
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _texto_data(valor):
    texto = _texto_celula(valor)
    if not texto:
        return ""
    formatos = (
        "%d/%m/%Y",
        "%d/%m/%y",
        "%Y-%m-%d",
        "%d-%m-%Y",
        "%Y-%m-%d %H:%M:%S",
    )
    for formato in formatos:
        try:
            data = datetime.strptime(texto, formato)
            return data.strftime("%d/%m/%Y")
        except ValueError:
            continue
    return texto


def _normalizar_cpf(valor):
    return re.sub(r"\D", "", str(valor or ""))


def _normalizar_matricula(valor):
    return re.sub(r"\s+", "", _normalizar_texto(valor))


def _codigo_inteiro(valor):
    texto = _texto_celula(valor)
    if not texto:
        return None
    try:
        codigo = int(float(texto.replace(",", ".")))
    except (TypeError, ValueError):
        return None
    return codigo if codigo > 0 else None


def _normalizar_sexo(valor):
    texto = _normalizar_texto(valor)
    if texto in {"M", "MASC", "MASCULINO"}:
        return "MASCULINO"
    if texto in {"F", "FEM", "FEMININO"}:
        return "FEMININO"
    return _texto_celula(valor).upper()


def _linha_para_dados(valores, mapeamento):
    dados = {}
    for indice, campo in mapeamento.items():
        if indice >= len(valores):
            continue
        valor = valores[indice]
        if campo == "codigo":
            codigo = _codigo_inteiro(valor)
            dados[campo] = codigo if codigo is not None else ""
        elif campo == "nascimento":
            dados[campo] = _texto_data(valor)
        elif campo == "sexo":
            dados[campo] = _normalizar_sexo(valor)
        elif campo == "uf":
            dados[campo] = _texto_celula(valor).upper()
        else:
            dados[campo] = _texto_celula(valor)

    if dados.get("nome"):
        dados["nome"] = re.sub(r"\s+", " ", dados["nome"]).strip()
    return dados


def _avaliar_cabecalho(valores):
    mapeamento = {}
    reconhecidos = []
    for indice, valor in enumerate(valores[:MAX_COLUNAS]):
        chave = _normalizar_cabecalho(valor)
        campo = ALIASES.get(chave)
        if campo and campo not in reconhecidos:
            mapeamento[indice] = campo
            reconhecidos.append(campo)

    # O relatório XLS fornecido pela escola pode trazer duas colunas de
    # filiação entre NOME e NASCIMENTO sem rótulos. Quando a estrutura é
    # inequívoca, usamos essas posições como Pai e Mãe.
    indices_por_campo = {campo: indice for indice, campo in mapeamento.items()}
    nome_idx = indices_por_campo.get("nome")
    nasc_idx = indices_por_campo.get("nascimento")
    if (
        nome_idx is not None
        and nasc_idx is not None
        and nasc_idx - nome_idx == 3
        and nome_idx + 1 not in mapeamento
        and nome_idx + 2 not in mapeamento
        and "pai" not in reconhecidos
        and "mae" not in reconhecidos
    ):
        mapeamento[nome_idx + 1] = "pai"
        mapeamento[nome_idx + 2] = "mae"
        reconhecidos.extend(["pai", "mae"])

    score = len(set(reconhecidos))
    if "nome" in reconhecidos:
        score += 4
    if "nascimento" in reconhecidos:
        score += 2
    return score, mapeamento


def _validar_xlsx_compactado(caminho):
    try:
        with zipfile.ZipFile(caminho) as arquivo:
            membros = arquivo.infolist()
            if len(membros) > MAX_XLSX_ARQUIVOS_INTERNOS:
                raise ErroImportacao(
                    "A planilha possui estrutura interna excessivamente grande."
                )

            total = sum(item.file_size for item in membros)
            if total > MAX_XLSX_DESCOMPACTADO:
                raise ErroImportacao(
                    "A planilha excede o limite seguro após descompactação."
                )

            for item in membros:
                if item.file_size > 25 * 1024 * 1024:
                    raise ErroImportacao(
                        "A planilha contém uma parte interna grande demais."
                    )
                if item.compress_size and item.file_size / item.compress_size > 200:
                    raise ErroImportacao(
                        "A planilha apresenta taxa de compressão insegura."
                    )
    except zipfile.BadZipFile as exc:
        raise ErroImportacao("O arquivo XLSX está corrompido.") from exc


def _linhas_xlsx(caminho):
    from openpyxl import load_workbook

    _validar_xlsx_compactado(caminho)
    wb = load_workbook(caminho, read_only=True, data_only=True)
    try:
        for planilha in wb.worksheets:
            linhas = []
            for numero, valores in enumerate(
                planilha.iter_rows(values_only=True),
                start=1,
            ):
                if numero > MAX_LINHAS + 40:
                    break
                linhas.append(list(valores[:MAX_COLUNAS]))
            yield planilha.title, linhas
    finally:
        wb.close()


def _linhas_xls(caminho):
    import xlrd

    wb = xlrd.open_workbook(
        filename=str(caminho),
        on_demand=True,
    )
    try:
        for planilha in wb.sheets():
            linhas = []
            limite = min(planilha.nrows, MAX_LINHAS + 40)
            for r in range(limite):
                valores = []
                for c in range(min(planilha.ncols, MAX_COLUNAS)):
                    cell = planilha.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            valor = xlrd.xldate.xldate_as_datetime(
                                cell.value,
                                wb.datemode,
                            )
                        except (ValueError, OverflowError):
                            valor = cell.value
                    else:
                        valor = cell.value
                    valores.append(valor)
                linhas.append(valores)
            yield planilha.name, linhas
    finally:
        wb.release_resources()


def ler_lista_alunos(caminho):
    caminho = Path(caminho)
    sufixo = caminho.suffix.lower()
    if sufixo == ".xlsx":
        fontes = _linhas_xlsx(caminho)
    elif sufixo == ".xls":
        fontes = _linhas_xls(caminho)
    else:
        raise ErroImportacao("Formato não suportado. Use .xls ou .xlsx.")

    melhor = None
    encontrou_lista_forte = False
    for nome_planilha, linhas in fontes:
        for indice, valores in enumerate(linhas[:40]):
            score, mapeamento = _avaliar_cabecalho(valores)
            if melhor is None or score > melhor["score"]:
                melhor = {
                    "score": score,
                    "planilha": nome_planilha,
                    "linhas": linhas,
                    "cabecalho_indice": indice,
                    "mapeamento": mapeamento,
                    "cabecalho": valores,
                }

            # Ao encontrar uma lista claramente identificada (como a aba
            # DADOS ALUNOS do arquivo oficial), para imediatamente. Isso evita
            # carregar dezenas de abas de atas e deixa o upload muito mais rápido.
            if score >= 10 and "nome" in mapeamento.values():
                encontrou_lista_forte = True
                break

        if encontrou_lista_forte:
            break

    if not melhor or melhor["score"] < 5 or "nome" not in melhor["mapeamento"].values():
        raise ErroImportacao(
            "Não foi possível identificar a linha de cabeçalho. "
            "A lista precisa conter pelo menos a coluna NOME e outros dados do aluno."
        )

    linhas_dados = []
    for numero_planilha, valores in enumerate(
        melhor["linhas"][melhor["cabecalho_indice"] + 1:],
        start=melhor["cabecalho_indice"] + 2,
    ):
        dados = _linha_para_dados(valores, melhor["mapeamento"])
        if not any(str(valor or "").strip() for valor in dados.values()):
            continue
        if not dados.get("nome"):
            continue
        dados["_linha"] = numero_planilha
        linhas_dados.append(dados)
        if len(linhas_dados) >= MAX_LINHAS:
            break

    if not linhas_dados:
        raise ErroImportacao("Nenhum aluno válido foi encontrado na lista.")

    reconhecidos = sorted(
        set(melhor["mapeamento"].values()),
        key=lambda campo: list(ROTULOS).index(campo)
        if campo in ROTULOS
        else 999,
    )
    usados = set(melhor["mapeamento"])
    ignorados = [
        _texto_celula(valor)
        for indice, valor in enumerate(melhor["cabecalho"])
        if indice not in usados and _texto_celula(valor)
    ]

    return {
        "planilha": melhor["planilha"],
        "linha_cabecalho": melhor["cabecalho_indice"] + 1,
        "campos": reconhecidos,
        "campos_rotulos": [ROTULOS[campo] for campo in reconhecidos],
        "ignorados": ignorados,
        "linhas": linhas_dados,
    }


def _adicionar_aos_indices(indices, aluno):
    indices["por_codigo"][aluno.codigo] = aluno

    cpf = _normalizar_cpf(aluno.cpf)
    if cpf and aluno not in indices["por_cpf"][cpf]:
        indices["por_cpf"][cpf].append(aluno)

    matricula = _normalizar_matricula(aluno.matricula)
    if matricula and aluno not in indices["por_matricula"][matricula]:
        indices["por_matricula"][matricula].append(aluno)

    nome = _normalizar_texto(aluno.nome)
    nascimento = _texto_data(aluno.nascimento)
    if nome:
        if aluno not in indices["por_nome"][nome]:
            indices["por_nome"][nome].append(aluno)
        if nascimento and aluno not in indices["por_nome_nascimento"][(nome, nascimento)]:
            indices["por_nome_nascimento"][(nome, nascimento)].append(aluno)


def _indices_alunos():
    indices = {
        "por_codigo": {},
        "por_cpf": defaultdict(list),
        "por_matricula": defaultdict(list),
        "por_nome_nascimento": defaultdict(list),
        "por_nome": defaultdict(list),
    }

    for aluno in Aluno.objects.all().order_by("id"):
        _adicionar_aos_indices(indices, aluno)

    return indices


def _unico(lista):
    if len(lista) == 1:
        return lista[0], None
    if len(lista) > 1:
        return None, "Mais de um aluno do banco corresponde a esta linha."
    return None, None


def localizar_aluno(dados, indices):
    codigo = dados.get("codigo")
    if codigo:
        aluno = indices["por_codigo"].get(codigo)
        if aluno:
            return aluno, "Código", None

    cpf = _normalizar_cpf(dados.get("cpf"))
    if cpf:
        aluno, erro = _unico(indices["por_cpf"].get(cpf, []))
        if aluno or erro:
            return aluno, "CPF", erro

    matricula = _normalizar_matricula(dados.get("matricula"))
    if matricula:
        aluno, erro = _unico(indices["por_matricula"].get(matricula, []))
        if aluno or erro:
            return aluno, "Matrícula", erro

    nome = _normalizar_texto(dados.get("nome"))
    nascimento = _texto_data(dados.get("nascimento"))
    if nome and nascimento:
        aluno, erro = _unico(
            indices["por_nome_nascimento"].get((nome, nascimento), [])
        )
        if aluno or erro:
            return aluno, "Nome + nascimento", erro

    if nome:
        aluno, erro = _unico(indices["por_nome"].get(nome, []))
        if aluno or erro:
            return aluno, "Nome", erro

    return None, "", None


def analisar_lista(resultado_leitura):
    indices = _indices_alunos()
    itens = []
    novos = 0
    existentes = 0
    conflitos = 0

    for dados in resultado_leitura["linhas"]:
        aluno, criterio, erro = localizar_aluno(dados, indices)
        if erro:
            status = "conflito"
            conflitos += 1
        elif aluno:
            status = "atualizar"
            existentes += 1
        else:
            status = "novo"
            novos += 1

        itens.append(
            {
                "linha": dados["_linha"],
                "nome": dados.get("nome", ""),
                "nascimento": dados.get("nascimento", ""),
                "codigo_lista": dados.get("codigo", ""),
                "status": status,
                "criterio": criterio,
                "aluno_codigo": aluno.codigo if aluno else "",
                "erro": erro or "",
            }
        )

    return {
        "total": len(itens),
        "novos": novos,
        "existentes": existentes,
        "conflitos": conflitos,
        "itens": itens,
    }


def _valor_limitado(campo, valor):
    field = Aluno._meta.get_field(campo)
    texto = _texto_celula(valor)
    if field.max_length:
        texto = texto[: field.max_length]
    return texto


@transaction.atomic
def aplicar_lista(resultado_leitura):
    indices = _indices_alunos()
    proximo_codigo = (
        Aluno.objects.aggregate(maximo=Max("codigo"))["maximo"] or 0
    ) + 1

    criados = 0
    atualizados = 0
    sem_alteracao = 0
    conflitos = []
    codigos_usados = set(indices["por_codigo"])

    for dados in resultado_leitura["linhas"]:
        aluno, criterio, erro = localizar_aluno(dados, indices)
        if erro:
            conflitos.append(
                {
                    "linha": dados["_linha"],
                    "nome": dados.get("nome", ""),
                    "erro": erro,
                }
            )
            continue

        if aluno is None:
            codigo_lista = dados.get("codigo")
            if codigo_lista and codigo_lista not in codigos_usados:
                codigo = codigo_lista
            else:
                while proximo_codigo in codigos_usados:
                    proximo_codigo += 1
                codigo = proximo_codigo
                proximo_codigo += 1

            valores = {
                "codigo": codigo,
                "nome": _valor_limitado("nome", dados.get("nome", "")),
                "curso": _valor_limitado(
                    "curso",
                    dados.get("curso") or "ENSINO FUNDAMENTAL",
                ),
                "ativo": True,
            }
            for campo in CAMPOS_ATUALIZAVEIS:
                if campo in {"nome", "curso"}:
                    continue
                valor = dados.get(campo)
                if valor not in (None, ""):
                    valores[campo] = _valor_limitado(campo, valor)

            aluno = Aluno.objects.create(**valores)
            criados += 1
            codigos_usados.add(aluno.codigo)
        else:
            alterados = []
            for campo in CAMPOS_ATUALIZAVEIS:
                if campo not in dados:
                    continue
                valor = dados.get(campo)
                # Células vazias da lista nunca apagam dados já existentes.
                if valor in (None, ""):
                    continue
                novo_valor = _valor_limitado(campo, valor)
                if getattr(aluno, campo) != novo_valor:
                    setattr(aluno, campo, novo_valor)
                    alterados.append(campo)

            if alterados:
                aluno.save(update_fields=alterados)
                atualizados += 1
            else:
                sem_alteracao += 1

        # Atualiza os índices em memória sem reler todo o banco a cada linha.
        # Mantemos também as chaves anteriores de um aluno atualizado, o que
        # ajuda a reconhecer linhas duplicadas do mesmo arquivo.
        _adicionar_aos_indices(indices, aluno)

    return {
        "total": len(resultado_leitura["linhas"]),
        "criados": criados,
        "atualizados": atualizados,
        "sem_alteracao": sem_alteracao,
        "conflitos": conflitos,
        "mantidos_fora_da_lista": True,
    }
