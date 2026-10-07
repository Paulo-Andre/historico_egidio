from pathlib import Path
from django.core.management.base import BaseCommand
from django.db import transaction
from openpyxl import load_workbook
from historico.models import Aluno, ParametroAno, AbaExcel, CelulaExcel, RegistroAta

class Command(BaseCommand):
    help='Importa integralmente o arquivo EMITIR HISTORICO.xlsx para o banco web.'
    def add_arguments(self, parser):
        parser.add_argument('arquivo', nargs='?', default='data/EMITIR HISTORICO.xlsx')

    @transaction.atomic
    def handle(self,*args,**opts):
        path=Path(opts['arquivo'])
        if not path.exists():
            raise SystemExit(f'Arquivo não encontrado: {path}')
        self.stdout.write('Lendo valores calculados...')
        wbv=load_workbook(path,data_only=True,read_only=False)
        self.stdout.write('Lendo fórmulas/estrutura...')
        wbf=load_workbook(path,data_only=False,read_only=False)
        Aluno.objects.all().delete()
        ParametroAno.objects.all().delete()
        RegistroAta.objects.all().delete()
        CelulaExcel.objects.all().delete()
        AbaExcel.objects.all().delete()

        ws=wbv['DADOS ALUNOS']
        for row in range(8,ws.max_row+1):
            codigo=ws.cell(row,1).value
            nome=ws.cell(row,2).value
            if not codigo or not nome:
                continue
            try:
                codigo=int(float(codigo))
            except:
                continue
            Aluno.objects.update_or_create(
                codigo=codigo,
                defaults={
                    'nome':str(nome).strip(),
                    'uf':str(ws.cell(row,3).value or '').strip(),
                    'pai':str(ws.cell(row,4).value or '').strip(),
                    'mae':str(ws.cell(row,5).value or '').strip(),
                    'nascimento':str(ws.cell(row,6).value or '').strip(),
                    'naturalidade':str(ws.cell(row,7).value or '').strip(),
                    'nacionalidade':str(ws.cell(row,8).value or '').strip(),
                    'sexo':str(ws.cell(row,9).value or '').strip(),
                })

        ws=wbv['DADOS ADCIONAIS']
        for row in range(8,ws.max_row+1):
            ano=ws.cell(row,1).value
            try:
                ano=int(ano)
            except:
                continue
            ParametroAno.objects.update_or_create(
                ano=ano,
                defaults={
                    'ch_ingles':str(ws.cell(row,2).value or ''),
                    'ch_anual':str(ws.cell(row,3).value or ''),
                    'dias_letivos':str(ws.cell(row,4).value or ''),
                    'ch_sem_ingles':str(ws.cell(row,5).value or ''),
                    'media':str(ws.cell(row,6).value or ''),
                    'escola':str(ws.cell(row,7).value or ''),
                })

        for ordem,nome in enumerate(wbf.sheetnames,1):
            wsf=wbf[nome]
            wsv=wbv[nome]
            aba=AbaExcel.objects.create(
                nome=nome,
                ordem=ordem,
                visivel=(wsf.sheet_state=='visible'),
                max_linha=wsf.max_row,
                max_coluna=wsf.max_column,
            )
            lote=[]
            for row in wsf.iter_rows():
                for cf in row:
                    vf=wsv[cf.coordinate].value
                    formula=cf.value if cf.data_type=='f' else ''
                    valor=vf if vf is not None else ('' if formula else cf.value)
                    if valor is None and not formula:
                        continue
                    if valor=='' and not formula:
                        continue
                    lote.append(CelulaExcel(
                        aba=aba,
                        referencia=cf.coordinate,
                        linha=cf.row,
                        coluna=cf.column,
                        valor=str(valor if valor is not None else ''),
                        formula=str(formula or ''),
                        numero_formato=str(cf.number_format or ''),
                        estilo_id=int(cf.style_id or 0),
                    ))
                    if len(lote)>=3000:
                        CelulaExcel.objects.bulk_create(lote)
                        lote=[]
            if lote:
                CelulaExcel.objects.bulk_create(lote)

            if nome.startswith('ATA '):
                try:
                    ano=int(nome.split()[-1])
                except:
                    continue
                registros=[]
                for r in range(1,wsf.max_row+1):
                    dados={}
                    for c in range(1,min(wsf.max_column,60)+1):
                        v=wsv.cell(r,c).value
                        if v not in (None,''):
                            dados[wsv.cell(r,c).column_letter]=str(v)
                    if dados:
                        texto=' | '.join(dados.values())
                        registros.append(RegistroAta(
                            ano=ano,
                            linha=r,
                            dados=dados,
                            texto_busca=texto[:12000],
                        ))
                        if len(registros)>=1000:
                            RegistroAta.objects.bulk_create(registros)
                            registros=[]
                if registros:
                    RegistroAta.objects.bulk_create(registros)
            self.stdout.write(f'  {nome}: importada')

        self.stdout.write(self.style.SUCCESS(
            f'Concluído: {Aluno.objects.count()} alunos, '
            f'{AbaExcel.objects.count()} abas e '
            f'{CelulaExcel.objects.count()} células preservadas.'
        ))
